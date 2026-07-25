import pandas as pd
from pathlib import Path
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Glosario explicativo de columnas
COLUMN_DESCRIPTIONS = [
    ("ticker", "Símbolo bursátil de la empresa en Wall Street (ej. AAPL, MSFT)."),
    ("sector", "Categoría o segmento de industria al que pertenece la empresa."),
    ("company_name", "Nombre legal o comercial completo de la compañía."),
    ("current_price", "Precio actual de cotización de la acción en dólares (USD)."),
    ("trailing_eps", "Earnings Per Share (EPS): Ganancia por acción de los últimos 12 meses."),
    ("forward_eps", "Ganancia por acción estimada/proyectada para los próximos 12 meses."),
    ("earnings_growth", "Tasa de crecimiento anual esperada de las ganancias."),
    ("book_value", "Valor de libro por acción (Patrimonio Neto / Total de acciones)."),
    ("free_cash_flow", "Flujo de caja libre total generado por la compañía en USD."),
    ("shares_outstanding", "Número total de acciones en circulación emitidas por la empresa."),
    ("intrinsic_value_graham", "Valor intrínseco teórico calculado mediante la fórmula de Benjamin Graham."),
    ("margin_of_safety_%", "Margen de seguridad (% de descuento respecto al valor intrínseco)."),
    ("aaa_rate_used", "Tasa de rendimiento de bonos corporativos AAA usada como tasa de descuento.")
]


def export_to_excel(
    comafi_df: pd.DataFrame,
    valuation_df: pd.DataFrame,
    filename: str = "cedear_valuation_report.xlsx",
    output_dir: str = "reports"
) -> Path:
    """Exporta DataFrames a Excel dentro del directorio especificado
    y retorna la ruta (Path) del archivo generado.
    """
    reports_path = Path(output_dir)
    reports_path.mkdir(parents=True, exist_ok=True)
    filepath = reports_path / filename

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        # 1. Escribir DataFrames
        valuation_df.to_excel(writer, sheet_name="Valuation", index=False)
        comafi_df.to_excel(writer, sheet_name="Comafi Ratios", index=False)
        
        # Estilos
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
        section_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # Formato base
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]

            for col_num in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
                for cell in row:
                    cell.border = thin_border
                    if isinstance(cell.value, str):
                        cell.alignment = left_align
                    else:
                        cell.alignment = right_align

        # Formatos numéricos en 'Valuation'
        val_ws = writer.sheets["Valuation"]
        col_indices = {cell.value: idx + 1 for idx, cell in enumerate(val_ws[1])}

        pct_cols = ["earnings_growth", "margin_of_safety_%", "aaa_rate_used"]
        for col_name in pct_cols:
            if col_name in col_indices:
                col_idx = col_indices[col_name]
                for row in range(2, len(valuation_df) + 2):
                    cell = val_ws.cell(row=row, column=col_idx)
                    cell.number_format = '0.00%'
                    if col_name in ["margin_of_safety_%", "aaa_rate_used"] and isinstance(cell.value, (int, float)):
                        cell.value = cell.value / 100.0

        num_cols = ["free_cash_flow", "shares_outstanding"]
        for col_name in num_cols:
            if col_name in col_indices:
                col_idx = col_indices[col_name]
                for row in range(2, len(valuation_df) + 2):
                    cell = val_ws.cell(row=row, column=col_idx)
                    cell.number_format = '#,##0'

        currency_cols = ["current_price", "trailing_eps", "forward_eps", "book_value", "intrinsic_value_graham"]
        for col_name in currency_cols:
            if col_name in col_indices:
                col_idx = col_indices[col_name]
                for row in range(2, len(valuation_df) + 2):
                    cell = val_ws.cell(row=row, column=col_idx)
                    cell.number_format = '#,##0.00'

        # Agregar Glosario
        start_row = len(valuation_df) + 4

        title_cell = val_ws.cell(row=start_row, column=1, value="Glosario y Descripción de Columnas")
        title_cell.font = section_font

        val_ws.cell(row=start_row + 1, column=1, value="Columna").font = bold_font
        val_ws.cell(row=start_row + 1, column=2, value="Descripción").font = bold_font

        for idx, (col_name, desc) in enumerate(COLUMN_DESCRIPTIONS, start=start_row + 2):
            cell_col = val_ws.cell(row=idx, column=1, value=col_name)
            cell_desc = val_ws.cell(row=idx, column=2, value=desc)
            
            cell_col.font = bold_font
            cell_desc.font = regular_font
            
            cell_col.border = thin_border
            cell_desc.border = thin_border

        # Autoajuste de ancho
        for sheet_name in writer.sheets:
            worksheet = writer.sheets[sheet_name]
            for col in worksheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                
                for cell in col:
                    val_str = str(cell.value or "")
                    if cell.row >= start_row and col_letter == "B" and sheet_name == "Valuation":
                        continue
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                
                worksheet.column_dimensions[col_letter].width = max(max_len + 5, 14)

    print(f"Successfully created '{filepath}' with formatted numbers and glossary!")
    return filepath